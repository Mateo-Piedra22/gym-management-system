import sys
import json
from pathlib import Path
import openpyxl

def list_placeholders(xlsx_path: str) -> list[str]:
    p = Path(xlsx_path)
    wb = openpyxl.load_workbook(str(p), data_only=False)
    placeholders = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and '{{' in v and '}}' in v:
                    placeholders.add(v.strip())
    try:
        wb.close()
    except Exception:
        pass
    return sorted(placeholders)

def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "usage: python scripts/list_template_placeholders.py <xlsx_template_path>"}))
        return
    xlsx_path = sys.argv[1]
    try:
        phs = list_placeholders(xlsx_path)
        # Limit output to avoid spam
        print(json.dumps({"count": len(phs), "placeholders": phs[:200]}, ensure_ascii=False, indent=2))
    except Exception as e:
        print(json.dumps({"error": str(e)}))

if __name__ == "__main__":
    main()